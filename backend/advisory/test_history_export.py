import csv
import datetime
import io

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from .history import EXPORT_COLUMNS
from .models import ActualWeatherReading, FarmProfile
from .test_history import add_readings


class HistoryExportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.user_a = User.objects.create_user("farmer_a", password="pw12345!")
        cls.user_b = User.objects.create_user("farmer_b", password="pw12345!")
        cls.staff = User.objects.create_user("admin_user", password="pw12345!", is_staff=True)
        today = datetime.date.today()
        cls.farm_a = FarmProfile.objects.create(farm_name="A Field", farmer_name="Farmer A", planting_date=today, owner=cls.user_a)
        cls.farm_b = FarmProfile.objects.create(farm_name="B Field", farmer_name="Farmer B", planting_date=today, owner=cls.user_b)
        cls.latest = today - datetime.timedelta(days=2)
        add_readings(cls.farm_a, 35, cls.latest)
        add_readings(cls.farm_b, 5, cls.latest)

    def download(self, query):
        return self.client.get(f"/api/history/export/{query}")

    def csv_rows(self, response):
        text = response.content.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text)))

    def test_csv_matches_the_screen_for_each_range(self):
        self.client.force_authenticate(self.user_a)
        for key, expected in (("last20", 20), ("7d", 7), ("30d", 30)):
            api = self.client.get(f"/api/history/?range={key}").data
            resp = self.download(f"?range={key}&file=csv")
            self.assertEqual(resp.status_code, 200)
            rows = self.csv_rows(resp)
            self.assertEqual(len(rows) - 1, expected)
            self.assertEqual([r["date"] for r in api["results"]], [r[0] for r in rows[1:]])

    def test_csv_headers_filename_and_excel_friendly_encoding(self):
        self.client.force_authenticate(self.user_a)
        resp = self.download("?file=csv")
        self.assertTrue(resp["Content-Type"].startswith("text/csv"))
        self.assertEqual(
            resp["Content-Disposition"],
            f'attachment; filename="agriaura-history-farm{self.farm_a.id}-last20.csv"',
        )
        self.assertTrue(resp.content.startswith(b"\xef\xbb\xbf"))  # BOM so Excel shows °C
        self.assertEqual(self.csv_rows(resp)[0], [h for h, _ in EXPORT_COLUMNS])

    def test_xlsx_has_data_sheet_and_info_sheet(self):
        self.client.force_authenticate(self.user_a)
        resp = self.download("?range=7d&file=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["History", "Info"])
        ws = wb["History"]
        self.assertEqual(ws.max_row, 8)  # heading + 7 days
        self.assertEqual(ws["A1"].value, "Date")
        self.assertEqual(ws["B2"].value, 30.0)
        info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
        self.assertEqual(info["Farmer"], "Farmer A")
        self.assertEqual(info["Rows"], 7)

    def test_formula_like_text_is_defused(self):
        ActualWeatherReading.objects.filter(farm=self.farm_a).update(condition_text='=HYPERLINK("http://evil")')
        self.client.force_authenticate(self.user_a)
        csv_cell = self.csv_rows(self.download("?file=csv"))[1][11]
        self.assertTrue(csv_cell.startswith("'="))
        xlsx = load_workbook(io.BytesIO(self.download("?file=xlsx").content))["History"]["L2"].value
        self.assertTrue(xlsx.startswith("'="))

    def test_access_rules_apply_to_downloads_too(self):
        self.assertEqual(self.download("?file=csv").status_code, 401)
        self.client.force_authenticate(self.user_b)
        self.assertEqual(self.download(f"?farm={self.farm_a.id}&file=csv").status_code, 404)
        self.assertEqual(len(self.csv_rows(self.download("?file=csv"))) - 1, 5)
        self.client.force_authenticate(self.staff)
        self.assertEqual(len(self.csv_rows(self.download(f"?farm={self.farm_a.id}&range=30d&file=csv"))) - 1, 30)

    def test_bad_parameters_are_400(self):
        self.client.force_authenticate(self.user_a)
        self.assertEqual(self.download("?range=forever").status_code, 400)
        self.assertEqual(self.download("?file=pdf").status_code, 400)
