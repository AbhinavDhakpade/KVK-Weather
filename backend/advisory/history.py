"""
History of a farm's observed weather: which rows to show for each time range.

Shared by the history list endpoint and (next step) the CSV/Excel download, so
what the person sees on screen is exactly what they download.

Ranges:
  last20  the 20 most recent daily readings
  7d      the 7 most recent days of data
  30d     the 30 most recent days of data

"Days" are counted back from the newest stored reading, not from today.
NASA POWER publishes observed weather with a delay of about 2 days, so
counting from today would always come up a couple of rows short.
"""

import csv
import datetime
import io

from django.utils import timezone

from .models import ActualWeatherReading

HISTORY_RANGES = {"last20": 20, "7d": 7, "30d": 30}


def history_rows(farm_id, range_key):
    """Readings for one farm in the given range, newest first."""
    qs = ActualWeatherReading.objects.filter(farm_id=farm_id).order_by("-date")
    if range_key == "last20":
        return list(qs[:20])

    latest = qs.values_list("date", flat=True).first()
    if latest is None:
        return []
    cutoff = latest - datetime.timedelta(days=HISTORY_RANGES[range_key] - 1)
    return list(qs.filter(date__gte=cutoff))


# ----------------------------------------------------------------- download ---

# (column heading, model field). Same order in the CSV and the Excel file.
EXPORT_COLUMNS = [
    ("Date", "date"),
    ("Max temp (°C)", "temp_max_c"),
    ("Min temp (°C)", "temp_min_c"),
    ("Humidity (%)", "humidity_pct"),
    ("Rainfall (mm)", "rainfall_mm"),
    ("ET0 (mm)", "et0_mm"),
    ("VPD (kPa)", "vpd_kpa"),
    ("Solar radiation (MJ/m²)", "solar_mj_m2"),
    ("Wind (km/h)", "wind_kmh"),
    ("GDD (daily)", "gdd_daily"),
    ("GDD (cumulative)", "gdd_cumulative"),
    ("Condition", "condition_text"),
    ("Source", "data_source"),
    ("Fetched at", "fetched_at"),
]

RANGE_LABELS = {
    "last20": "20 most recent daily readings",
    "7d": "7 most recent days",
    "30d": "30 most recent days",
}


def _safe_text(value):
    """Stop spreadsheet programs from running text that starts like a formula."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def _local_naive(moment):
    return timezone.localtime(moment).replace(tzinfo=None) if moment else None


def csv_bytes(rows):
    """UTF-8 with a byte-order mark, so Excel shows °C and ² correctly."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow([heading for heading, _ in EXPORT_COLUMNS])
    for row in rows:
        out = []
        for _, field in EXPORT_COLUMNS:
            value = getattr(row, field)
            if field == "fetched_at":
                local = _local_naive(value)
                value = local.strftime("%Y-%m-%d %H:%M") if local else ""
            elif isinstance(value, str):
                value = _safe_text(value)
            out.append(value)
        writer.writerow(out)
    return buf.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows, farm, range_key):
    """One sheet of data (headings in row 1, ready for Excel/pandas) plus an Info sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "History"
    ws.append([heading for heading, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        values = []
        for _, field in EXPORT_COLUMNS:
            value = getattr(row, field)
            if field == "fetched_at":
                value = _local_naive(value)
            elif isinstance(value, str):
                value = _safe_text(value)
            values.append(value)
        ws.append(values)

    ws.freeze_panes = "A2"
    for idx, (heading, field) in enumerate(EXPORT_COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = max(12, len(heading) + 2)
        if field == "date":
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        elif field == "fetched_at":
            ws.column_dimensions[letter].width = 18
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm"

    info = wb.create_sheet("Info")
    for label, value in [
        ("Farm", farm.farm_name),
        ("Farmer", farm.farmer_name),
        ("Range", RANGE_LABELS[range_key]),
        ("Rows", len(rows)),
        ("Data sources", ", ".join(sorted({r.data_source for r in rows})) or "none"),
        ("Order", "Newest first"),
        ("Exported at", timezone.localtime().replace(tzinfo=None).strftime("%Y-%m-%d %H:%M")),
    ]:
        info.append([label, _safe_text(value) if isinstance(value, str) else value])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 56
    for cell in info["A"]:
        cell.font = Font(bold=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
